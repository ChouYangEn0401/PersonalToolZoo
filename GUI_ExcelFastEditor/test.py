import tkinter as tk
from tkinter import ttk, messagebox, Menu
import pandas as pd
from openpyxl import Workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter


class DataFrameViewerApp:
    def __init__(self, root, data):
        self.root = root
        self.root.title("動態資料檢視器與輸出器")
        self.root.geometry("1200x600")

        self.data = data
        self.processed_df = self.data.copy()
        self.group_columns = ['學校', '部門', '小指標']  # 預設索引欄位
        self.pivot_column = '大指標'  # 預設欄位合併欄位

        self.create_widgets()
        self.create_menu()
        self.update_tree()

    def create_widgets(self):
        # Treeview
        tree_frame = tk.Frame(self.root)
        tree_frame.pack(fill='both', expand=True, padx=10, pady=10)

        tree_scroll_y = ttk.Scrollbar(tree_frame, orient='vertical')
        tree_scroll_y.pack(side='right', fill='y')
        tree_scroll_x = ttk.Scrollbar(tree_frame, orient='horizontal')
        tree_scroll_x.pack(side='bottom', fill='x')

        self.tree = ttk.Treeview(tree_frame, yscrollcommand=tree_scroll_y.set, xscrollcommand=tree_scroll_x.set)
        self.tree.pack(fill='both', expand=True)
        tree_scroll_y.config(command=self.tree.yview)
        tree_scroll_x.config(command=self.tree.xview)

        # 按鈕區
        button_frame = tk.Frame(self.root)
        button_frame.pack(pady=10)

        view_frame = ttk.LabelFrame(button_frame, text="視圖切換")
        view_frame.pack(side='left', padx=10, pady=5)
        ttk.Button(view_frame, text="顯示原始資料", command=lambda: self.show_view('raw')).pack(side='left', padx=5,
                                                                                                pady=5)
        ttk.Button(view_frame, text="列合併視圖", command=lambda: self.show_view('row')).pack(side='left', padx=5,
                                                                                              pady=5)
        ttk.Button(view_frame, text="欄合併視圖", command=lambda: self.show_view('column')).pack(side='left', padx=5,
                                                                                                 pady=5)

        export_frame = ttk.LabelFrame(button_frame, text="檔案輸出")
        export_frame.pack(side='left', padx=10, pady=5)
        ttk.Button(export_frame, text="輸出原始資料", command=lambda: self.export_view('raw')).pack(side='left', padx=5,
                                                                                                    pady=5)
        ttk.Button(export_frame, text="輸出列合併", command=lambda: self.export_view('row')).pack(side='left', padx=5,
                                                                                                  pady=5)
        ttk.Button(export_frame, text="輸出欄合併", command=lambda: self.export_view('column')).pack(side='left',
                                                                                                     padx=5, pady=5)
        ttk.Button(export_frame, text="輸出複合合併", command=lambda: self.export_view('combined')).pack(side='left',
                                                                                                         padx=5, pady=5)

        # 狀態列
        self.status_bar = tk.Label(self.root, text="就緒", bd=1, relief=tk.SUNKEN, anchor=tk.W)
        self.status_bar.pack(side='bottom', fill='x')

    def create_menu(self):
        menubar = Menu(self.root)
        self.root.config(menu=menubar)

        options_menu = Menu(menubar, tearoff=0)
        menubar.add_cascade(label="設定", menu=options_menu)

        # 建立索引欄位選單
        index_menu = Menu(options_menu, tearoff=0)
        options_menu.add_cascade(label="設定索引欄位 (Row Merge)", menu=index_menu)
        for col in self.data.columns:
            index_menu.add_command(label=col, command=lambda c=col: self.set_group_column(c))

        # 建立欄位合併欄位選單
        pivot_menu = Menu(options_menu, tearoff=0)
        options_menu.add_cascade(label="設定欄位合併欄位 (Column Merge)", menu=pivot_menu)
        for col in self.data.columns:
            pivot_menu.add_command(label=col, command=lambda c=col: self.set_pivot_column(c))

    def set_group_column(self, col_name):
        self.group_columns = [col_name]
        messagebox.showinfo("設定成功", f"已將 '{col_name}' 設定為索引欄位。")
        self.status_bar.config(text=f"索引欄位已設定為: {col_name}")

    def set_pivot_column(self, col_name):
        self.pivot_column = col_name
        messagebox.showinfo("設定成功", f"已將 '{col_name}' 設定為欄位合併欄位。")
        self.status_bar.config(text=f"欄位合併欄位已設定為: {col_name}")

    def update_tree(self, df=None, add_header=None):
        """更新 Treeview 顯示內容的通用方法"""
        self.tree.delete(*self.tree.get_children())
        if df is None:
            df = self.processed_df

        cols = list(df.columns)
        self.tree["columns"] = cols
        self.tree["show"] = "headings"

        # 設定欄位標頭
        for col in cols:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=120, anchor='center')

        # 處理額外的群組標頭
        if add_header:
            self.tree.insert("", "end", values=add_header, tags=('header_row',))
            self.tree.tag_configure('header_row', background='#E0E0E0', font=('Arial', 10, 'bold'))

        # 插入資料
        for _, row in df.iterrows():
            self.tree.insert("", "end", values=list(row.fillna('')))

    def show_view(self, view_type):
        """根據類型切換並顯示視圖"""
        self.status_bar.config(text="正在處理資料...")
        try:
            if view_type == 'raw':
                self.processed_df = self.data.copy()
                self.update_tree()
            elif view_type == 'row':
                self.processed_df = self.data.copy()
                for col in self.group_columns:
                    self.processed_df[col] = self.processed_df[col].mask(self.processed_df[col].duplicated())
                self.update_tree()
            elif view_type == 'column':
                if not self.pivot_column or self.pivot_column not in self.data.columns:
                    messagebox.showwarning("警告", "請先透過 '設定' 選單選擇一個欄位合併欄位。")
                    return

                try:
                    pivot_df = self.data.pivot_table(
                        index=self.group_columns,
                        columns=self.pivot_column,
                        values='分數',
                        aggfunc='first'
                    ).reset_index()
                    self.processed_df = pivot_df.copy()

                    group_header = self.group_columns + [self.pivot_column] * (
                                len(pivot_df.columns) - len(self.group_columns))
                    self.update_tree(df=self.processed_df, add_header=group_header)

                except KeyError as e:
                    messagebox.showerror("錯誤", f"在進行欄位合併時發生錯誤，請檢查所選的欄位是否包含在資料中: {e}")
                    return

            self.status_bar.config(text=f"已載入 '{view_type}' 視圖。")
        except Exception as e:
            messagebox.showerror("錯誤", f"處理視圖時發生錯誤: {e}")
            self.status_bar.config(text="操作失敗。")

    def _merge_cells_in_excel(self, ws, df_for_merge, group_cols, start_row_offset):
        """處理 Excel 儲存格合併的輔助方法"""
        for col_name in group_cols:
            col_idx = df_for_merge.columns.get_loc(col_name) + 1  # get_loc 回傳 0-based，openpyxl 是 1-based
            current_value = None
            start_row = None

            for idx, value in enumerate(df_for_merge[col_name], start=start_row_offset):
                if value != current_value:
                    if start_row and idx - start_row > 1:
                        ws.merge_cells(start_row=start_row, start_column=col_idx, end_row=idx - 1, end_column=col_idx)
                        ws.cell(start_row, col_idx).alignment = Alignment(vertical="center", horizontal="center")
                    current_value = value
                    start_row = idx

            # 合併最後一組
            if start_row and len(df_for_merge) + start_row_offset - start_row > 0:
                ws.merge_cells(start_row=start_row, start_column=col_idx,
                               end_row=len(df_for_merge) + start_row_offset - 1, end_column=col_idx)
                ws.cell(start_row, col_idx).alignment = Alignment(vertical="center", horizontal="center")

    def export_view(self, view_type):
        """根據類型輸出檔案"""
        file_path = f"output_{view_type}.xlsx"
        self.status_bar.config(text=f"正在輸出檔案: {file_path}")
        try:
            if view_type == 'raw':
                self.data.to_excel(file_path, index=False)
            elif view_type == 'row':
                wb = Workbook()
                ws = wb.active
                ws.title = "Row Merged"
                df = self.data.copy()
                ws.append(list(df.columns))
                for r in dataframe_to_rows(df, index=False, header=False):
                    ws.append(r)
                self._merge_cells_in_excel(ws, df, self.group_columns, start_row_offset=2)
                wb.save(file_path)
            elif view_type in ['column', 'combined']:
                if not self.pivot_column or self.pivot_column not in self.data.columns:
                    messagebox.showwarning("警告", "請先透過 '設定' 選單選擇一個欄位合併欄位。")
                    return

                pivot = self.data.pivot_table(
                    index=self.group_columns,
                    columns=self.pivot_column,
                    values='分數',
                    aggfunc='first'
                ).reset_index()

                wb = Workbook()
                ws = wb.active
                ws.title = "Column Merged" if view_type == 'column' else "Combined Merged"

                group_header = self.group_columns + [self.pivot_column] * (len(pivot.columns) - len(self.group_columns))
                ws.append(group_header)
                ws.append(list(pivot.columns))

                for r in dataframe_to_rows(pivot, index=False, header=False):
                    ws.append(r)

                # 欄合併
                start_col = len(self.group_columns) + 1
                if len(pivot.columns) > len(self.group_columns):
                    ws.merge_cells(start_row=1, start_column=start_col, end_row=1, end_column=len(pivot.columns))
                    ws.cell(1, start_col).value = self.pivot_column
                    ws.cell(1, start_col).alignment = Alignment(horizontal="center", vertical="center")

                for col in range(1, len(self.group_columns) + 1):
                    ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
                    ws.cell(1, col).value = self.group_columns[col - 1]
                    ws.cell(1, col).alignment = Alignment(horizontal="center", vertical="center")

                # 列合併 (僅針對 combined 模式)
                if view_type == 'combined':
                    self._merge_cells_in_excel(ws, pivot, self.group_columns, start_row_offset=3)

                wb.save(file_path)

            messagebox.showinfo("成功", f"已輸出 {file_path}")
            self.status_bar.config(text=f"檔案已成功輸出至: {file_path}")

        except Exception as e:
            messagebox.showerror("錯誤", f"輸出時發生錯誤: {e}")
            self.status_bar.config(text="操作失敗。")


if __name__ == '__main__':
    # 原始資料
    data = pd.DataFrame({
        '學校': ['A', 'A', 'A', 'B', 'B', 'B'],
        '部門': ['Math', 'English', 'English', 'Math', 'English', 'English'],
        '大指標': ['Educational', 'Educational', 'Research', 'Research', 'Educational', 'Exam'],
        '小指標': ['Text', 'Text', 'Text', 'Text', 'Text', 'Text'],
        '分數': [90, 100, 96, 88, 98, 98]
    })

    root = tk.Tk()
    app = DataFrameViewerApp(root, data)
    root.mainloop()