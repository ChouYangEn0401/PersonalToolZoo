import os
import fitz  # PyMuPDF


def pdf_to_text(input_path, output_path=None):
    """Extract all text from a PDF and save to a .txt file."""
    doc = fitz.open(input_path)
    text_parts = []
    for page_num in range(len(doc)):
        page = doc[page_num]
        text_parts.append(f"--- 第 {page_num + 1} 頁 ---\n")
        text_parts.append(page.get_text())
        text_parts.append("\n")
    doc.close()

    full_text = "".join(text_parts)

    if output_path is None:
        base = os.path.splitext(input_path)[0]
        output_path = base + ".txt"

    with open(output_path, "w", encoding="utf-8") as f:
        f.write(full_text)
    return output_path


def pdf_to_docx(input_path, output_path=None):
    """Convert PDF to Word document using pdf2docx."""
    from pdf2docx import Converter

    if output_path is None:
        base = os.path.splitext(input_path)[0]
        output_path = base + ".docx"

    cv = Converter(input_path)
    cv.convert(output_path)
    cv.close()
    return output_path


def pdf_to_pptx(input_path, output_path=None):
    """Convert PDF to PPTX by rendering each page as an image slide."""
    from pptx import Presentation
    from pptx.util import Inches
    import io

    if output_path is None:
        base = os.path.splitext(input_path)[0]
        output_path = base + ".pptx"

    doc = fitz.open(input_path)
    prs = Presentation()

    for page_num in range(len(doc)):
        page = doc[page_num]
        # Render page to image at 200 DPI
        mat = fitz.Matrix(200 / 72, 200 / 72)
        pix = page.get_pixmap(matrix=mat)
        img_bytes = pix.tobytes("png")

        # Set slide dimensions to match page aspect ratio
        page_width = page.rect.width
        page_height = page.rect.height
        aspect = page_width / page_height

        slide_width = Inches(10)
        slide_height = Inches(10 / aspect)
        prs.slide_width = slide_width
        prs.slide_height = slide_height

        slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
        img_stream = io.BytesIO(img_bytes)
        slide.shapes.add_picture(img_stream, 0, 0, slide_width, slide_height)

    doc.close()
    prs.save(output_path)
    return output_path
